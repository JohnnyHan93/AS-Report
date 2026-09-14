from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from as_report.analyzer import analyze
from as_report.cleaner import clean_data
from as_report.input_set import (
    AS_RAW_ROLE,
    CUSTOMER_ROBOT_MASTER_ROLE,
    FAILURE_PART_MASTER_ROLE,
    ApprovedInputSet,
    InputFileIdentity,
)
from as_report.input_set_validation import build_input_set_validation, normalize_robot_model_key
from as_report.period import PeriodSelection, build_period_range, filter_by_period
from as_report.report_excel import write_excel_report
from as_report.report_html import write_html_report
from conftest import sample_dataframe


def _identity(role: str, path: Path, rows: int) -> InputFileIdentity:
    return InputFileIdentity(
        role=role,
        path=path,
        file_name=path.name,
        size=100,
        modified_time_ns=1,
        sha256=f"hash-{role}",
        row_count=rows,
        column_count=4,
        unique_id_count=rows,
        duplicate_id_count=0,
        columns=("*ID",),
        date_min="2025-01-01" if role == AS_RAW_ROLE else "",
        date_max="2025-12-31" if role == AS_RAW_ROLE else "",
    )


def _input_set(tmp_path: Path, raw_rows: int) -> ApprovedInputSet:
    return ApprovedInputSet(
        input_set_id="TEST-INPUT-SET",
        manifest_path=tmp_path / "active_input_set.json",
        manifest_sha256="manifest-hash",
        files={
            AS_RAW_ROLE: _identity(AS_RAW_ROLE, tmp_path / "raw.csv", raw_rows),
            CUSTOMER_ROBOT_MASTER_ROLE: _identity(CUSTOMER_ROBOT_MASTER_ROLE, tmp_path / "robot.csv", 4),
            FAILURE_PART_MASTER_ROLE: _identity(FAILURE_PART_MASTER_ROLE, tmp_path / "part.csv", 3),
        },
        canonical_mappings={"robot_model_raw": {"status": "approved"}},
    )


def _robot_master() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"*ID": "M1", "본체(Manifold)": "MPX3500", "고객사(Customer)": "A"},
            {"*ID": "M2", "본체(Manifold)": "MPX 3500", "고객사(Customer)": "A"},
            {"*ID": "M3", "본체(Manifold)": "Fanuc", "고객사(Customer)": "B"},
            {"*ID": "M4", "본체(Manifold)": "PAINT", "고객사(Customer)": "B"},
        ]
    )


def _part_master() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"*ID": "P1", "품명": "AOPR", "대분류": "Spray", "중분류": "Applicator"},
            {"*ID": "P2", "품명": "AOPR", "대분류": "Spray", "중분류": "PCF"},
            {"*ID": "P3", "품명": "PartA", "대분류": "Robot", "중분류": "Controller"},
        ]
    )


def test_robot_model_normalization_is_safe_and_limited() -> None:
    assert normalize_robot_model_key("  mh50Ⅱ ‐ 20 ") == "MH50II-20"
    assert normalize_robot_model_key("MH50(2)-20") == "MH50(2)-20"
    assert normalize_robot_model_key("MH50(2)-20") != normalize_robot_model_key("MH50-2")
    assert normalize_robot_model_key("선택안함") == ""


def test_gate_b_robot_model_validation_separates_non_model_values(tmp_path: Path) -> None:
    raw = pd.DataFrame(
        [
            {"*ID": "A1", "접수일": "2025-01-01", "고객사": "A", "로보트 기종": "MPX 3500", "소분류 (고장부품)": "PartA", "대분류": "Robot", "중분류": "Controller"},
            {"*ID": "A2", "접수일": "2025-01-02", "고객사": "B", "로보트 기종": "FANUC", "소분류 (고장부품)": "AOPR", "대분류": "Spray", "중분류": "Applicator"},
            {"*ID": "A3", "접수일": "2025-01-03", "고객사": "B", "로보트 기종": "MH50(2)-20", "소분류 (고장부품)": "Unknown", "대분류": "Robot", "중분류": "Arm"},
            {"*ID": "A4", "접수일": "2025-01-04", "고객사": "B", "로보트 기종": "", "소분류 (고장부품)": "AOPR, PartA", "대분류": "Spray", "중분류": "Applicator"},
        ]
    )
    result = build_input_set_validation(
        raw,
        _robot_master(),
        _part_master(),
        _input_set(tmp_path, len(raw)),
        run_id="RUN-1",
        period_type="year",
        period_label="2025년",
        start_date="2025-01-01",
        end_date="2025-12-31",
    )

    match_values = result["robot_model_match_summary"].set_index("항목")["값"]
    assert match_values["설치자산 기종 안전 키 매칭"] == 2
    assert match_values["canonical 적용 가능 AS 건수"] == 1
    assert match_values["검토 필요 AS 건수"] == 2

    install = result["robot_model_install_summary"].set_index("robot_model_key")
    assert install.loc["MPX3500", "등록 설치대수"] == 2
    assert install.loc["MPX3500", "매칭된 AS 접수건수"] == 1
    assert "FANUC" not in install.index
    assert set(result["robot_model_unmatched_preview"]["robot_model_raw"]) == {"FANUC", "MH50(2)-20"}

    part_values = result["part_quality_summary"].set_index("항목")["값"]
    assert part_values["고장부품 입력"] == 4
    assert part_values["품명 exact match"] == 2
    assert part_values["AOPR 분류 충돌"] == 1
    assert part_values["multi-value"] == 1


def test_report_generation_ignores_inactive_input_set_validation_payload(tmp_path: Path) -> None:
    raw = sample_dataframe()
    robot_master = pd.DataFrame(
        [
            {"*ID": "M1", "본체(Manifold)": "ModelA", "고객사(Customer)": "고객A"},
            {"*ID": "M2", "본체(Manifold)": "ModelB", "고객사(Customer)": "고객B"},
        ]
    )
    part_master = pd.DataFrame(
        [
            {"*ID": "P1", "품명": "부품A", "대분류": "대분류A", "중분류": "중분류A"},
            {"*ID": "P2", "품명": "부품C", "대분류": "대분류C", "중분류": "중분류C"},
        ]
    )
    validation = build_input_set_validation(
        raw,
        robot_master,
        part_master,
        _input_set(tmp_path, len(raw)),
        run_id="RUN-SAME-123",
        period_type="year",
        period_label="2025년",
        start_date="2025-01-01",
        end_date="2025-12-31",
        created_at="2026-07-22T10:00:00",
    )

    cleaned = clean_data(raw)
    period_range = build_period_range(PeriodSelection(period="year", year=2025))
    period_data = filter_by_period(cleaned, period_range)
    analysis = analyze(cleaned, period_data, period_range=period_range, comparison_source_data=cleaned, source_period_data=period_data)
    analysis["input_set_validation"] = validation
    html_path = tmp_path / "gate_b.html"
    excel_path = tmp_path / "gate_b.xlsx"
    write_html_report(
        html_path,
        analysis,
        ["Gate B validation"],
        period_label=period_range.label,
        start_date=period_range.start_date.isoformat(),
        end_date=period_range.end_date.isoformat(),
        source_file_name="raw.csv",
    )
    write_excel_report(excel_path, analysis)

    html = html_path.read_text(encoding="utf-8")
    assert "AS_REPORT_LAYOUT_VERSION: AS17_BRIEFING" in html
    assert "승인 입력 세트 Gate B 검증" not in html
    assert "RUN-SAME-123" not in html
    assert "등록 설치대수 기준 대당 AS 접수건수" not in html

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    assert "21_Input_Set_Validation" not in workbook.sheetnames
    assert "22_Robot_Model_Validation" not in workbook.sheetnames
    assert "23_Part_Master_Validation" not in workbook.sheetnames
