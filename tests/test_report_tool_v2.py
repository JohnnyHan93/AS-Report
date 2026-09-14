from dataclasses import replace
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from pptx import Presentation
import pandas as pd
import pytest

from as_report.analyzer import analyze
from as_report.cleaner import clean_data
from as_report.period import PeriodSelection, build_period_range, filter_by_period
from as_report.report_payload import ReportPayload, build_report_payload
from as_report.report_qc import check_report, check_artifacts
from as_report.report_tool import generate_structured_documents
from conftest import sample_dataframe


@pytest.fixture
def payload():
    data = clean_data(sample_dataframe())
    period = build_period_range(PeriodSelection(period="year", year=2025))
    analysis = analyze(data, filter_by_period(data, period), period_range=period)
    provenance = {"files": {"as_raw": {"path": "C:/private/raw.csv", "file_name": "raw.csv", "sha256": "a" * 64, "row_count": 3, "unique_id_count": 3, "date_min": "2024-12-31", "date_max": "2025-07-20"}}}
    return build_report_payload(analysis, ["선택 기간의 접수 현황입니다."], period_range=period, filters={}, provenance=provenance, run_id="test-v2", owner="tester", generated_at="2026-09-11T12:00:00", analysis_version="test")


def test_payload_roundtrip_and_provenance(payload):
    assert ReportPayload.from_json(payload.to_json()) == payload
    assert payload.to_json() == payload.to_json()
    assert "C:/private" not in payload.to_json()
    assert not [i for i in check_report(payload) if i.severity == "error"]
    assert payload.metadata["master_identity"] is None


def test_renderers_and_same_run_identity(payload):
    outputs = generate_structured_documents(payload)
    doc = Document(BytesIO(outputs["report.docx"]))
    ppt = Presentation(BytesIO(outputs["report.pptx"]))
    assert len(doc.tables) >= 3
    assert all(s.title in [p.text for p in doc.paragraphs] for s in payload.sections)
    assert len(ppt.slides) >= 8
    assert not check_artifacts(payload, {k: v for k, v in outputs.items() if k.endswith(("docx", "pptx"))})
    assert "ppt/slideMasters/slideMaster1.xml" in ZipFile(BytesIO(outputs["report.pptx"])).namelist()
    assert ReportPayload.from_json(outputs["report_payload.json"].decode()) == payload


@pytest.mark.parametrize("key", ["run_id", "owner", "analysis_version", "generated_at", "raw_identity"])
def test_missing_metadata_blocks_generation(payload, key):
    broken = replace(payload, metadata={k: v for k, v in payload.metadata.items() if k != key})
    with pytest.raises(ValueError, match="문서 검수 실패"):
        generate_structured_documents(broken)


def test_authored_prose_and_raw_values_distinguished(payload):
    sections = list(payload.sections)
    sections[0] = replace(sections[0], paragraphs=["협력사 책임 소재 확인"])
    assert any(i.code == "unsafe_narrative" for i in check_report(replace(payload, sections=sections)))
    table = sections[1].tables[0]
    sections[1] = replace(sections[1], tables=[replace(table, rows=[["책임님", 2, "건"]])])
    sections[0] = payload.sections[0]
    assert not any(i.code == "unsafe_narrative" for i in check_report(replace(payload, sections=sections)))


@pytest.mark.parametrize("value", [float("nan"), float("inf"), "NaN", "None", None])
def test_bad_metrics_block(payload, value):
    section = payload.sections[1]
    table = replace(section.tables[0], rows=[["접수", value, "건"]])
    broken = replace(payload, sections=[payload.sections[0], replace(section, tables=[table]), *payload.sections[2:]])
    assert any(i.severity == "error" for i in check_report(broken))


def test_mismatched_pair(payload):
    outputs = generate_structured_documents(payload, ("docx",))
    different = replace(payload, scope={**payload.scope, "filters": {"고객사": ["다른 고객사"]}})
    assert check_artifacts(different, {"report.docx": outputs["report.docx"]})


def test_formats_are_independent(payload, monkeypatch):
    import as_report.report_pptx as ppt
    monkeypatch.setattr(ppt, "render_pptx", lambda _: pytest.fail("PPT should not run"))
    assert "report.docx" in generate_structured_documents(payload, ("docx",))


def test_missing_sections_and_evidence(payload):
    assert any(i.code == "missing_sections" for i in check_report(replace(payload, sections=[])))
    section = replace(payload.sections[0], evidence_status="invalid")
    assert any(i.code == "invalid_evidence" for i in check_report(replace(payload, sections=[section, *payload.sections[1:]])))


def test_invalid_chart_and_percentage(payload):
    section = replace(payload.sections[1], tables=[])
    assert any(i.code == "missing_metric" for i in check_report(replace(payload, sections=[payload.sections[0], section, *payload.sections[2:]])))
    bad = replace(payload, charts=[{"title": "missing", "section": "kpis", "table": "not present", "x": "x", "y": "y"}])
    assert any(i.code == "invalid_chart" for i in check_report(bad))
    section = payload.sections[1]
    table = replace(section.tables[0], rows=[["대당 AS 접수건수", 10, "%"]])
    assert any(i.code == "invalid_percentage" for i in check_report(replace(payload, sections=[payload.sections[0], replace(section, tables=[table]), *payload.sections[2:]])))


def test_streamlit_current_run_downloads(payload, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from types import SimpleNamespace
    import as_report.report_payload as model
    import as_report.report_tool as tool
    monkeypatch.setattr(model, "build_report_payload", lambda *args, **kwargs: payload)
    monkeypatch.setattr(tool, "generate_structured_documents", lambda *args: {"report.docx": b"test"})
    def screen():
        import streamlit as st
        from types import SimpleNamespace
        from as_report.app import render_structured_report_actions
        result = SimpleNamespace(context_signature=st.session_state.get("test_context", "first"), run_timestamp="20260911_120000_000001", analysis={}, narrative_lines=[], period_range=None, filters={}, input_provenance={}, output_owner="tester")
        render_structured_report_actions(result)
    app = AppTest.from_function(screen).run()
    app.button(key="prepare_v2_documents").click().run()
    assert not app.exception
    assert len(app.get("download_button")) == 1
    app.run()
    assert len(app.get("download_button")) == 1
    app.session_state["test_context"] = "second"
    app.run()
    assert not app.get("download_button")
