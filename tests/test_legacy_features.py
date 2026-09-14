from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from as_report.analyzer import analyze
from as_report.cleaner import clean_data
from as_report.legacy_features import build_legacy_features
from as_report.master import build_master_analysis_result
from as_report.period import PeriodSelection, build_period_range, filter_by_period
from as_report.report_excel import write_excel_report
from as_report.report_html import write_html_report


def _legacy_sample() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "*ID": 1,
                "상태": "처리완료",
                "접수일": "2026-01-10",
                "업무유형": "A/S대응(부품수리/클레임처리)",
                "유/무상": "",
                "제조사": "",
                "고객사": "A",
                "BOOTH": "B1",
                "LINE": "L1",
                "공정": "도장",
                "ZONE": "Z1",
                "로보트 NO": "R1",
                "로보트 기종": "Robot-A",
                "대분류": "로보트",
                "중분류": "컨트롤러",
                "소분류 (고장부품)": "OS카드",
                "라인 중단 시간 (분)": 10,
                "접수 내용 (요약)": "접수 원문",
                "처리내용/진행상황 (요약)": "조치 원문",
            },
            {
                "*ID": 2,
                "상태": "전달 완료",
                "접수일": "2026-01-20",
                "업무유형": "A/S대응(일반 방문)",
                "유/무상": "유상",
                "제조사": "Maker",
                "고객사": "A",
                "BOOTH": "B1",
                "LINE": "L2",
                "공정": "도장",
                "ZONE": "Z1",
                "로보트 NO": "R2",
                "로보트 기종": "Robot-B",
                "대분류": "기타",
                "중분류": "기타",
                "소분류 (고장부품)": "기타",
                "라인 중단 시간 (분)": 0,
                "접수 내용 (요약)": "",
                "처리내용/진행상황 (요약)": "",
            },
            {
                "*ID": 3,
                "상태": "진행중",
                "접수일": "2026-02-05",
                "업무유형": "A/S대응(부품수리/클레임처리)",
                "유/무상": "무상",
                "제조사": "Maker",
                "고객사": "B",
                "BOOTH": "B2",
                "LINE": "L1",
                "공정": "도장",
                "ZONE": "Z2",
                "로보트 NO": "R3",
                "로보트 기종": "Robot-C",
                "대분류": "도장기",
                "중분류": "EVO",
                "소분류 (고장부품)": "EVO",
                "라인 중단 시간 (분)": 0,
                "접수 내용 (요약)": "클레임 접수",
                "처리내용/진행상황 (요약)": "확인 중",
            },
        ]
    )


def _robot_master(include_controller: bool = True, include_robot_model: bool = False) -> pd.DataFrame:
    rows = [
        {
            "*ID": "M1",
            "설치연도(Installation year)": "2025",
            "고객사(Customer)": "A",
            "위치(Location)": "화성",
            "공장(plant)": "1공장",
            "BOOTH": "B1",
            "LINE": "L1",
            "ZONE": "Z1",
            "Robot No": "R1",
            "공정-대분류 (major category)": "도장",
            "공정-중분류 (middle category)": "도장",
            "컨트롤러(Controller)": "YRC1000",
        },
        {
            "*ID": "M2",
            "설치연도(Installation year)": "2024",
            "고객사(Customer)": "A",
            "위치(Location)": "화성",
            "공장(plant)": "1공장",
            "BOOTH": "B1",
            "LINE": "L2",
            "ZONE": "Z1",
            "Robot No": "R2",
            "공정-대분류 (major category)": "도장",
            "공정-중분류 (middle category)": "도장",
            "컨트롤러(Controller)": "YRC1000",
        },
    ]
    dataframe = pd.DataFrame(rows)
    if include_robot_model:
        dataframe["Robot Model"] = ["AR1440", "AR1440"]
    if not include_controller:
        dataframe = dataframe.drop(columns=["컨트롤러(Controller)"])
    return dataframe


def _failure_master() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"*ID": "P1", "제조사": "Maker", "품명": "OS카드", "중분류": "컨트롤러", "대분류": "로보트", "검색용": "OS카드"},
            {"*ID": "P2", "제조사": "Maker", "품명": "EVO", "중분류": "EVO", "대분류": "도장기", "검색용": "EVO"},
        ]
    )


def test_claim_analysis_uses_missing_label_for_manufacturer_and_paid_type() -> None:
    cleaned = clean_data(_legacy_sample())
    result = build_legacy_features(cleaned)

    assert result["claim_summary"].loc[0, "값"] == 2
    assert result["claim_note"] == "업무유형 입력값 기준 검토용입니다. 입력값을 검토할 뿐 추가 판단을 확정하지 않습니다."
    assert "미입력" in result["claim_by_manufacturer"]["제조사"].tolist()
    assert "미입력" in result["claim_by_paid_type"]["유/무상"].tolist()
    assert {"접수내용", "조치이력"}.issubset(result["claim_issue_list"].columns)


def test_completion_status_kpi_counts_delivered_done_only_in_as25_output() -> None:
    cleaned = clean_data(_legacy_sample())
    result = build_legacy_features(cleaned)
    kpi = result["completion_kpi"].set_index("항목")["값"].to_dict()

    assert kpi["분석 대상 접수 건수"] == 3
    assert kpi["처리완료 건수"] == 2
    assert kpi["진행/미완료 건수"] == 1


def test_missing_status_column_returns_safe_limited_result() -> None:
    dataframe = clean_data(_legacy_sample()).drop(columns=["상태", "처리상태구분"])
    result = build_legacy_features(dataframe)

    assert "상태 컬럼" in result["completion_note"]
    assert result["completion_kpi"].iloc[-1]["값"] == "상태 컬럼 없음"


def test_controller_preview_uses_controller_column() -> None:
    cleaned = clean_data(_legacy_sample())
    master_result = build_master_analysis_result(cleaned, _robot_master(), _failure_master())

    assert "YRC1000" in master_result["controller_install_summary"]["컨트롤러"].tolist()
    row = master_result["controller_install_summary"].set_index("컨트롤러").loc["YRC1000"]
    assert row["설치대수"] == 2
    assert row["AS 접수 건수"] == 2


def test_missing_controller_column_does_not_substitute_robot_model() -> None:
    cleaned = clean_data(_legacy_sample())
    master_result = build_master_analysis_result(cleaned, _robot_master(include_controller=False), _failure_master())

    assert master_result["controller_note"] == "컨트롤러 기준 컬럼 없음"
    assert master_result["controller_install_summary"].empty


def test_robot_model_preview_uses_only_explicit_robot_model_column() -> None:
    cleaned = clean_data(_legacy_sample())
    master_result = build_master_analysis_result(
        cleaned,
        _robot_master(include_robot_model=True),
        _failure_master(),
    )

    row = master_result["robot_model_install_summary"].set_index("로보트 기종").loc["AR1440"]
    assert row["설치대수"] == 2
    assert row["AS 접수 건수"] == 2
    assert row["대당 AS 접수건수"] == 1.0
    assert "로봇 master 매칭률(%)" in master_result["robot_model_match_summary"]["항목"].tolist()


def test_robot_model_preview_is_limited_without_robot_model_column() -> None:
    cleaned = clean_data(_legacy_sample())
    master_result = build_master_analysis_result(cleaned, _robot_master(), _failure_master())

    assert master_result["robot_model_install_summary"].empty
    assert "로보트 기종 기준 설치대수 컬럼" in master_result["robot_model_note"]


def test_raw_only_html_and_excel_exclude_master_sections(tmp_path: Path) -> None:
    cleaned = clean_data(_legacy_sample())
    period_range = build_period_range(PeriodSelection(period="year", year=2026))
    period_data = filter_by_period(cleaned, period_range)
    analysis = analyze(cleaned, period_data, period_range=period_range, comparison_source_data=cleaned, source_period_data=period_data)
    html_path = tmp_path / "validation.html"
    excel_path = tmp_path / "validation_summary.xlsx"
    write_html_report(
        html_path,
        analysis,
        ["AS25 validation"],
        period_label=period_range.label,
        start_date=period_range.start_date.isoformat(),
        end_date=period_range.end_date.isoformat(),
        source_file_name="sample.csv",
    )
    write_excel_report(excel_path, analysis)

    html = html_path.read_text(encoding="utf-8")
    assert "처리 상태 요약" in html
    assert "클레임/제조사 검토" in html
    assert "로보트 기종별 등록 설치대수 기준 대당 AS 접수건수" not in html
    assert "로봇 master 매칭 현황" not in html
    assert "승인 입력 세트 Gate B 검증" not in html
    assert "컨트롤러별 설치대수" not in html
    assert "업무유형 입력값 기준 검토용입니다. 입력값을 검토할 뿐 추가 판단을 확정하지 않습니다." in html
    assert "원인이나 책임을 확정하지 않습니다" not in html
    assert "책임" not in html
    assert "고장률" not in html

    workbook = load_workbook(excel_path, read_only=True)
    assert {"17_Claim_Analysis", "19_Completion_Status"}.issubset(workbook.sheetnames)
    assert "18_Robot_Model_AS" not in workbook.sheetnames
    assert "21_Input_Set_Validation" not in workbook.sheetnames
    assert "22_Robot_Model_Validation" not in workbook.sheetnames
    assert "23_Part_Master_Validation" not in workbook.sheetnames
    assert "18_Controller_AS_Rate" not in workbook.sheetnames
