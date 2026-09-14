from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from as_report.analyzer import analyze
from as_report.cleaner import clean_data
from as_report.period import PeriodSelection, build_period_range, filter_by_period
from as_report.quality_feedback import build_quality_feedback_candidates
from as_report.report_excel import write_excel_report
from as_report.report_html import write_html_report


def _sample() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "*ID": 1,
                "상태": "처리완료",
                "접수일": "2026-01-10",
                "업무유형": "A/S대응(부품수리/클레임처리)",
                "유/무상": "",
                "제조사": "",
                "고객사": "A",
                "BOOTH": "B1",
                "LINE": "L1",
                "공정": "도장",
                "ZONE": "Z1",
                "로보트 NO": "R1",
                "로보트 기종": "Robot-A",
                "대분류": "로보트",
                "중분류": "컨트롤러",
                "소분류 (고장부품)": "OS카드",
                "라인 중단 시간 (분)": 10,
                "접수 내용 (요약)": "접수 원문",
                "처리내용/진행상황 (요약)": "조치 원문",
            },
            {
                "*ID": 2,
                "상태": "진행중",
                "접수일": "2026-01-20",
                "업무유형": "A/S대응(부품수리/클레임처리)",
                "유/무상": "무상",
                "제조사": "Maker",
                "고객사": "A",
                "BOOTH": "B1",
                "LINE": "L2",
                "공정": "도장",
                "ZONE": "Z1",
                "로보트 NO": "R2",
                "로보트 기종": "Robot-B",
                "대분류": "로보트",
                "중분류": "컨트롤러",
                "소분류 (고장부품)": "OS카드",
                "라인 중단 시간 (분)": 0,
                "접수 내용 (요약)": "클레임 접수",
                "처리내용/진행상황 (요약)": "확인 중",
            },
            {
                "*ID": 3,
                "상태": "처리완료",
                "접수일": "2026-02-05",
                "업무유형": "A/S대응(일반 방문)",
                "유/무상": "유상",
                "제조사": "Other",
                "고객사": "B",
                "BOOTH": "B2",
                "LINE": "L1",
                "공정": "도장",
                "ZONE": "Z2",
                "로보트 NO": "R3",
                "로보트 기종": "Robot-C",
                "대분류": "도장기",
                "중분류": "EVO",
                "소분류 (고장부품)": "EVO",
                "라인 중단 시간 (분)": 0,
                "접수 내용 (요약)": "일반 접수",
                "처리내용/진행상황 (요약)": "완료",
            },
        ]
    )


def test_quality_feedback_candidates_use_safe_input_based_groups() -> None:
    cleaned = clean_data(_sample())
    result = build_quality_feedback_candidates(cleaned)

    summary = result["quality_feedback_summary"].set_index("항목")["값"].to_dict()
    assert summary["클레임/제조사 입력 검토 후보"] == 2
    assert summary["유/무상/제조사 미입력 검토 후보"] == 1
    assert "미입력" in result["quality_feedback_claim_candidates"]["제조사"].tolist()
    assert "유/무상, 제조사" in result["quality_feedback_missing_field_candidates"]["미입력 항목"].tolist()
    assert result["quality_feedback_note"] == "입력값 기준으로 추가 확인이 필요한 후보를 모은 참고 섹션입니다. 표시 항목은 확정 판단이 아닙니다."


def test_quality_feedback_html_and_excel_outputs(tmp_path: Path) -> None:
    cleaned = clean_data(_sample())
    period_range = build_period_range(PeriodSelection(period="year", year=2026))
    period_data = filter_by_period(cleaned, period_range)
    analysis = analyze(cleaned, period_data, period_range=period_range, comparison_source_data=cleaned, source_period_data=period_data)

    html_path = tmp_path / "validation.html"
    excel_path = tmp_path / "validation_summary.xlsx"
    write_html_report(
        html_path,
        analysis,
        ["AS27 validation"],
        period_label=period_range.label,
        start_date=period_range.start_date.isoformat(),
        end_date=period_range.end_date.isoformat(),
        source_file_name="sample.csv",
    )
    write_excel_report(excel_path, analysis)

    html = html_path.read_text(encoding="utf-8")
    assert "품질 피드백 검토 후보" in html
    assert "처리 상태 요약" in html
    assert "클레임/제조사 검토" in html
    assert "로보트 기종별 등록 설치대수 기준 대당 AS 접수건수" not in html
    assert "고장률" not in html
    assert "책임" not in html
    assert "귀책" not in html
    assert "원인 확정" not in html

    workbook = load_workbook(excel_path, read_only=True)
    assert {"17_Claim_Analysis", "19_Completion_Status", "20_Quality_Feedback"}.issubset(workbook.sheetnames)
    assert "18_Robot_Model_AS" not in workbook.sheetnames
