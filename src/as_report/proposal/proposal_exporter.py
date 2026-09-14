from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from as_report.output import build_timestamped_filename, get_run_timestamp

from .proposal_engine import ProposalResult
from .proposal_rules import INSPECTION_RULES, PROPOSAL_DISCLAIMER


def export_proposal_outputs(result: ProposalResult, output_dir: str | Path, timestamp: str | None = None) -> dict[str, Path]:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    run_timestamp = timestamp or get_run_timestamp()
    paths = {
        "proposal_candidate_summary": output_path / build_timestamped_filename("proposal_candidate_summary", ".md", run_timestamp),
        "preventive_maintenance_candidates": output_path / build_timestamped_filename("preventive_maintenance_candidates", ".xlsx", run_timestamp),
        "horizontal_deployment_candidates": output_path / build_timestamped_filename("horizontal_deployment_candidates", ".xlsx", run_timestamp),
        "preventive_maintenance_candidates_csv": output_path / build_timestamped_filename("preventive_maintenance_candidates", ".csv", run_timestamp),
        "horizontal_deployment_candidates_csv": output_path / build_timestamped_filename("horizontal_deployment_candidates", ".csv", run_timestamp),
    }
    _write_summary(paths["proposal_candidate_summary"], result)
    _write_workbook(paths["preventive_maintenance_candidates"], "예방점검_후보", result.preventive_candidates)
    _write_workbook(paths["horizontal_deployment_candidates"], "수평전개_후보", result.horizontal_candidates)
    _to_csv(result.preventive_candidates, paths["preventive_maintenance_candidates_csv"])
    _to_csv(result.horizontal_candidates, paths["horizontal_deployment_candidates_csv"])
    paths.update(export_refined_proposal_outputs(result, output_path, run_timestamp))
    return paths


def export_refined_proposal_outputs(result: ProposalResult, output_dir: str | Path, timestamp: str | None = None) -> dict[str, Path]:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    run_timestamp = timestamp or get_run_timestamp()
    paths = {
        "proposal_top_candidates": output_path / build_timestamped_filename("proposal_top_candidates", ".xlsx", run_timestamp),
        "proposal_customer_summary": output_path / build_timestamped_filename("proposal_customer_summary", ".xlsx", run_timestamp),
        "proposal_candidate_summary_refined": output_path / build_timestamped_filename("proposal_candidate_summary_refined", ".md", run_timestamp),
    }
    _write_top_candidates_workbook(paths["proposal_top_candidates"], result)
    _write_customer_summary_workbook(paths["proposal_customer_summary"], result)
    _write_refined_summary(paths["proposal_candidate_summary_refined"], result)
    return paths


def _write_summary(path: Path, result: ProposalResult) -> None:
    preventive = result.preventive_candidates
    horizontal = result.horizontal_candidates
    summary = result.summary
    lines = [
        "# Proposal Candidate Summary",
        "",
        "## 분석 요약",
        f"- analysis period: {summary.get('analysis_period', '')}",
        f"- AS analysis row count: {summary.get('as_analysis_row_count', 0):,}",
        f"- preventive inspection candidate count: {len(preventive):,}",
        f"- preventive TOP candidate count: {summary.get('preventive_top_candidate_count', 0):,}",
        f"- horizontal deployment candidate count: {len(horizontal):,}",
        f"- horizontal TOP candidate count: {summary.get('horizontal_top_candidate_count', 0):,}",
        "",
        "## Candidate Counts By Inspection Item",
        _count_markdown(preventive, "inspection_item"),
        "",
        "## Candidate Counts By Customer",
        _count_markdown(preventive, "고객사"),
        "",
        "## Top Candidate Customers/Factories",
        _top_customer_factory(horizontal),
        "",
        "## Warnings And Limitations",
        "- 모든 결과는 제안 후보이며 담당자 검토 전 확정값으로 사용하지 않습니다.",
        "- 수평전개 후보는 exact/key 기반 매칭만 사용하며 fuzzy matching은 사용하지 않았습니다.",
        "- 설치연도 누락 건은 사용연차 기반 후보에서 제외하거나 manual_review로 표시했습니다.",
        "",
        "## Disclaimer",
        PROPOSAL_DISCLAIMER,
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_workbook(path: Path, sheet_name: str, dataframe: pd.DataFrame) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "데이터가 없습니다."}])
        output.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_workbook(writer)


def _write_top_candidates_workbook(path: Path, result: ProposalResult) -> None:
    preventive_top = _top_rows(result.preventive_candidates)
    horizontal_top = _top_rows(result.horizontal_candidates)
    matrix = _inspection_item_matrix()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _sheet(preventive_top, "예방점검_TOP").to_excel(writer, sheet_name="예방점검_TOP", index=False)
        _sheet(horizontal_top, "수평전개_TOP").to_excel(writer, sheet_name="수평전개_TOP", index=False)
        matrix.to_excel(writer, sheet_name="inspection_item_matrix", index=False)
        _format_workbook(writer)


def _write_customer_summary_workbook(path: Path, result: ProposalResult) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _customer_summary(result).to_excel(writer, sheet_name="customer_summary", index=False)
        _format_workbook(writer)


def _write_refined_summary(path: Path, result: ProposalResult) -> None:
    summary = result.summary
    preventive = result.preventive_candidates
    horizontal = result.horizontal_candidates
    lines = [
        "# Proposal Candidate Summary Refined",
        "",
        "## 후보 수 요약",
        f"- analysis period: {summary.get('analysis_period', '')}",
        f"- AS analysis row count: {summary.get('as_analysis_row_count', 0):,}",
        f"- preventive all candidate count: {summary.get('preventive_inspection_candidate_count', len(preventive)):,}",
        f"- preventive TOP candidate count: {summary.get('preventive_top_candidate_count', 0):,}",
        f"- horizontal all candidate count: {summary.get('horizontal_deployment_candidate_count', len(horizontal)):,}",
        f"- horizontal TOP candidate count: {summary.get('horizontal_top_candidate_count', 0):,}",
        f"- proposal TOP candidate count: {summary.get('proposal_top_candidate_count', 0):,}",
        "",
        "## TOP 후보 기준",
        "- 수평전개 High: same_model_same_install_year + 정보성 있는 source_part_category + target_factory/설치년도 존재",
        "- 수평전개 Medium: same_model_near_install_year + 정보성 있는 source_part_category + target_factory/설치년도 존재",
        "- same_customer_other_line, target_factory 없음, 설치년도 없음, source_part_category 비정보성 값은 TOP에서 제외하거나 Low로 표시했습니다.",
        "- 예방점검 후보는 inspection_item applicability matrix 기준으로 항목을 분리했습니다.",
        "",
        "## 예방점검 TOP 후보 By Inspection Item",
        _count_markdown(_top_rows(preventive), "inspection_item"),
        "",
        "## 수평전개 TOP 후보 By Customer",
        _count_markdown(_top_rows(horizontal), "target_customer"),
        "",
        "## Warnings And Limitations",
        "- 모든 결과는 제안 후보이며 담당자 검토 전 확정값으로 사용하지 않습니다.",
        "- fuzzy matching은 사용하지 않았습니다.",
        "- 비용 절감, 예방 효과, 대책 확정을 의미하지 않습니다.",
        "- 고객 제안 전 담당자 검토가 필요합니다.",
        "",
        "## Disclaimer",
        PROPOSAL_DISCLAIMER,
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _format_workbook(writer: pd.ExcelWriter) -> None:
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 24)
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = 10
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 60))
            worksheet.column_dimensions[column_letter].width = max_length + 2


def _to_csv(dataframe: pd.DataFrame, path: Path) -> None:
    output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "데이터가 없습니다."}])
    output.to_csv(path, index=False, encoding="utf-8-sig")


def _top_rows(dataframe: pd.DataFrame) -> pd.DataFrame:
    if dataframe.empty or "top_candidate" not in dataframe.columns:
        return pd.DataFrame(columns=dataframe.columns)
    return dataframe.loc[dataframe["top_candidate"].fillna(False).astype(bool)].copy()


def _sheet(dataframe: pd.DataFrame, empty_message: str) -> pd.DataFrame:
    return dataframe if not dataframe.empty else pd.DataFrame([{"안내": f"{empty_message} 데이터가 없습니다."}])


def _inspection_item_matrix() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "rule_id": rule.rule_id,
                "inspection_item": rule.inspection_item,
                "applicability_group": rule.applicability_group,
                "keywords": ", ".join(rule.keywords),
            }
            for rule in INSPECTION_RULES
        ]
    )


def _customer_summary(result: ProposalResult) -> pd.DataFrame:
    preventive = result.preventive_candidates.copy()
    horizontal = result.horizontal_candidates.copy()
    rows: list[dict[str, Any]] = []
    customers = set()
    if not preventive.empty and "고객사" in preventive.columns:
        customers.update(preventive["고객사"].fillna("").astype(str).tolist())
    if not horizontal.empty and "target_customer" in horizontal.columns:
        customers.update(horizontal["target_customer"].fillna("").astype(str).tolist())

    for customer in sorted(value for value in customers if value):
        preventive_rows = preventive.loc[preventive.get("고객사", pd.Series(index=preventive.index, dtype="object")).astype(str) == customer]
        horizontal_rows = horizontal.loc[horizontal.get("target_customer", pd.Series(index=horizontal.index, dtype="object")).astype(str) == customer]
        rows.append(
            {
                "customer": customer,
                "preventive_all_count": int(len(preventive_rows)),
                "preventive_top_count": int(preventive_rows.get("top_candidate", pd.Series(dtype=bool)).fillna(False).sum()) if not preventive_rows.empty else 0,
                "horizontal_all_count": int(len(horizontal_rows)),
                "horizontal_top_count": int(horizontal_rows.get("top_candidate", pd.Series(dtype=bool)).fillna(False).sum()) if not horizontal_rows.empty else 0,
                "high_horizontal_count": int((horizontal_rows.get("top_candidate_level", pd.Series(dtype="object")) == "High").sum()) if not horizontal_rows.empty else 0,
                "medium_horizontal_count": int((horizontal_rows.get("top_candidate_level", pd.Series(dtype="object")) == "Medium").sum()) if not horizontal_rows.empty else 0,
            }
        )
    return pd.DataFrame(rows) if rows else pd.DataFrame([{"안내": "데이터가 없습니다."}])


def _count_markdown(dataframe: pd.DataFrame, column: str) -> str:
    if dataframe.empty or column not in dataframe.columns:
        return "- 데이터가 없습니다."
    counts = dataframe[column].fillna("미입력").astype(str).value_counts().head(10).reset_index()
    counts.columns = [column, "count"]
    return _markdown_table(counts)


def _top_customer_factory(dataframe: pd.DataFrame) -> str:
    if dataframe.empty:
        return "- 데이터가 없습니다."
    working = dataframe.copy()
    working["customer_factory"] = working.get("target_customer", "").astype(str) + " / " + working.get("target_factory", "").astype(str)
    counts = working["customer_factory"].value_counts().head(10).reset_index()
    counts.columns = ["customer_factory", "count"]
    return _markdown_table(counts)


def _markdown_table(dataframe: pd.DataFrame) -> str:
    if dataframe.empty:
        return "- 데이터가 없습니다."
    columns = dataframe.columns.tolist()
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for _, row in dataframe.iterrows():
        lines.append("| " + " | ".join(_cell(row.get(column)) for column in columns) + " |")
    return "\n".join(lines)


def _cell(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).replace("\n", " ")
