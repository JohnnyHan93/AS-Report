"""Word renderer: consumes an approved payload and returns download bytes."""
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .report_payload import ReportPayload, display_value, report_style
from .report_qc import embed_identity, require_valid


def render_docx(payload: ReportPayload) -> bytes:
    require_valid(payload)
    style = report_style()
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(1.8)
    page.left_margin = page.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        font = doc.styles[name].font
        font.name = style["font"]
        doc.styles[name].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), style["font"])
    doc.styles["Normal"].font.size = Pt(style["docx_body_size"])
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_heading("A/S 분석 보고서", 0)
    doc.add_paragraph(f"조회 기간  {payload.scope['start']} ~ {payload.scope['end']}")
    doc.add_paragraph("선택한 입력 데이터의 접수 현황과 추가 확인 항목을 정리한 보고서입니다.")
    raw = payload.metadata["raw_identity"]
    doc.add_paragraph(f"입력 파일  {raw['file_name']}\n원본 범위  {raw['date_min']} ~ {raw['date_max']} / {raw['row_count']}건\n분석 ID  {payload.metadata['run_id']}\n생성일  {payload.metadata['generated_at']}")
    filters = "; ".join(f"{k}: {', '.join(v)}" for k, v in payload.scope["filters"].items() if v)
    doc.add_paragraph("선택 조건  " + (filters or "추가 필터 없음"))
    for section in payload.sections:
        doc.add_heading(section.title, 1)
        for text in section.paragraphs:
            doc.add_paragraph(text)
        for source_table in section.tables:
            doc.add_heading(source_table.title, 2)
            if not source_table.rows:
                doc.add_paragraph("표시할 데이터가 없습니다.")
                continue
            table = doc.add_table(rows=1, cols=len(source_table.columns))
            table.style = "Light Shading Accent 1"
            repeat_header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
            for cell, label in zip(table.rows[0].cells, source_table.columns):
                cell.text = label
            for row in source_table.rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = display_value(value)
        doc.add_paragraph("근거: " + {"Observed": "입력 데이터 집계", "Inferred": "검토 후보", "Unverified": "추가 확인 필요"}[section.evidence_status])
    doc.core_properties.identifier = payload.metadata["run_id"]
    doc.core_properties.title = "A/S 분석 보고서"
    output = BytesIO()
    doc.save(output)
    return embed_identity(output.getvalue(), payload)
