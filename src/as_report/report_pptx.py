"""Office-independent renderer using the repository-owned company master."""
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from .report_payload import ReportPayload, display_value, report_style
from .report_qc import embed_identity, require_valid


def render_pptx(payload: ReportPayload) -> bytes:
    require_valid(payload)
    style = report_style()
    presentation = Presentation(str(Path(__file__).parent / "resources/DY_PPT_Template_16x9.pptx"))
    layout = next((l for m in presentation.slide_masters for l in m.slide_layouts if l.name == style["ppt_layout"]), None)
    if layout is None:
        raise ValueError("회사 양식의 본문 레이아웃을 찾을 수 없습니다.")
    # Remove sample slides in memory; retain the approved masters and backgrounds.
    for slide_id in list(presentation.slides._sldIdLst):
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)

    def text(slide, value, x, y, w, h, size, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for index, line in enumerate(value.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = style["font"]
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(style["ink"])

    def new_slide(title):
        slide = presentation.slides.add_slide(layout)
        for shape in list(slide.placeholders):
            element = shape._element
            element.getparent().remove(element)
        text(slide, title, .55, .5, 12.1, .6, style["ppt_title_size"], True)
        slide.notes_slide.notes_text_frame.text = f"run_id: {payload.metadata['run_id']}\nRaw: {payload.metadata['raw_identity']['file_name']}\npayload: {payload.fingerprint}"
        return slide

    cover = new_slide("A/S 분석 보고서")
    text(cover, f"{payload.scope['start']} ~ {payload.scope['end']}\n\n{payload.metadata['raw_identity']['file_name']}\n생성일 {payload.metadata['generated_at']}", .8, 2, 11.5, 3, 22)
    index = new_slide("목차")
    text(index, "\n".join(f"{i}. {s.title}" for i, s in enumerate(payload.sections, 1)), .8, 1.5, 11.5, 5, 21)
    for section in payload.sections:
        for spec in [c for c in payload.charts if c["section"] == section.key]:
            source = next(t for t in section.tables if t.title == spec["table"])
            for offset in range(0, len(source.rows), 12):
                rows = source.rows[offset:offset + 12]
                data = CategoryChartData()
                data.categories = [display_value(row[source.columns.index(spec["x"])]) for row in rows]
                data.add_series(spec["y"], [row[source.columns.index(spec["y"])] for row in rows])
                slide = new_slide(spec["title"])
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.5), Inches(11.9), Inches(5.1), data).chart
                chart.has_legend = False
                chart.category_axis.tick_labels.font.name = style["font"]
                chart.category_axis.tick_labels.font.size = Pt(14)
                chart.value_axis.tick_labels.font.size = Pt(12)
                chart.series[0].format.fill.solid()
                chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(style["blue"])
        if section.paragraphs:
            for start in range(0, len(section.paragraphs), 4):
                slide = new_slide(section.title)
                lines = section.paragraphs[start:start + 4]
                if any(len(line) > 210 for line in lines):
                    raise ValueError(f"{section.title}: 문장이 길어 PPT 배치 검토가 필요합니다.")
                text(slide, "\n\n".join(lines), .65, 1.5, 12, 5.1, 20)
        for source_table in section.tables:
            if not source_table.rows:
                continue
            display_rows = source_table.rows[:style["ppt_table_limit"]]
            # Never silently clip long Raw strings. Reject with an actionable error.
            if any(len(display_value(v)) > 110 for row in display_rows for v in row):
                raise ValueError(f"{source_table.title}: 긴 원문은 Word 문서에서 확인해 주세요.")
            rows_per_slide = style["ppt_rows"]
            if any(len(display_value(v)) > 40 for row in display_rows for v in row):
                rows_per_slide = 4
            for offset in range(0, len(display_rows), rows_per_slide):
                rows = display_rows[offset:offset + rows_per_slide]
                limited = len(display_rows) < len(source_table.rows)
                title = source_table.title + (f" (처음 {len(display_rows)}행)" if limited else "")
                slide = new_slide(title)
                if limited:
                    slide.notes_slide.notes_text_frame.text += f"\n전체 {len(source_table.rows)}행. 전체 내용은 동일 실행 Word 문서 및 payload에서 확인할 수 있습니다."
                table = slide.shapes.add_table(len(rows) + 1, len(source_table.columns), Inches(.55), Inches(1.55), Inches(12.2), Inches(min(4.95, (len(rows) + 1) * .55))).table
                if "조건값" in source_table.columns and len(source_table.columns) == 5:
                    for column, width in zip(table.columns, style["repeat_column_widths"]):
                        column.width = Inches(width)
                    table.rows[0].height = Inches(.55)
                    for row in list(table.rows)[1:]:
                        row.height = Inches(4.4 / len(rows))
                for r, values in enumerate([source_table.columns, *rows]):
                    for c, value in enumerate(values):
                        cell = table.cell(r, c)
                        cell.text = display_value(value)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(style["blue"] if r == 0 else style["light"] if r % 2 else "FFFFFF")
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name = style["font"]
                            paragraph.font.size = Pt(style["ppt_body_size"])
                            paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else style["ink"])
                            paragraph.font.bold = r == 0
    presentation.core_properties.identifier = payload.metadata["run_id"]
    output = BytesIO()
    presentation.save(output)
    return embed_identity(output.getvalue(), payload)
