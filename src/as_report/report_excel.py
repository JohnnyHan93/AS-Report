from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .reporting.display_formatters import evaluate_repeat_condition_quality, format_repeat_condition_display

SHEETS = {
    "kpi": "00_KPI",
    "monthly": "01_Monthly_Trend",
    "work_type": "02_WorkType",
    "status": "03_Status",
    "customer": "04_Customer",
    "customer_downtime": "05_Customer_Downtime",
    "booth_line_process": "06_Booth_Line_Process",
    "robot_model": "07_Robot_Model",
    "cause": "08_Cause",
    "category": "09_Category",
    "part": "10_Part",
    "downtime_top": "11_Downtime_Top",
    "open_issues": "12_Open_Issues",
    "quality": "13_Data_Quality",
    "comparison": "14_Comparison",
    "repeat_issues": "15_Repeat_Issues",
    "raw_quality": "16_Raw_Quality_Check",
    "claim_analysis": "17_Claim_Analysis",
    "completion_status": "19_Completion_Status",
    "quality_feedback": "20_Quality_Feedback",
}


def write_excel_report(path: Path, analysis: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = analysis["tables"]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        write_table(writer, _dict_to_frame(analysis["kpi"]), SHEETS["kpi"])
        write_multi_tables(
            writer,
            SHEETS["monthly"],
            [("월별 접수 건수", tables["월별 접수 건수"]), ("월별 라인 중단 시간", tables["월별 라인 중단 시간"])],
        )
        write_table(writer, tables["업무유형구분별 접수 건수"], SHEETS["work_type"])
        write_table(writer, tables["상태별 접수 건수"], SHEETS["status"])
        write_table(writer, tables["고객사별 접수 건수 TOP 20"], SHEETS["customer"])
        write_table(writer, tables["고객사별 라인 중단 시간 TOP 20"], SHEETS["customer_downtime"])
        write_multi_tables(
            writer,
            SHEETS["booth_line_process"],
            [
                ("BOOTH별 접수 건수", tables["BOOTH별 접수 건수"]),
                ("LINE별 접수 건수", tables["LINE별 접수 건수"]),
                ("공정별 접수 건수", tables["공정별 접수 건수"]),
            ],
        )
        write_table(writer, tables["로보트 기종별 접수 건수"], SHEETS["robot_model"])
        write_table(writer, tables["고장원인별 접수 건수"], SHEETS["cause"])
        write_multi_tables(
            writer,
            SHEETS["category"],
            [
                ("대분류별 접수 건수", tables["대분류별 접수 건수"]),
                ("중분류별 접수 건수", tables["중분류별 접수 건수"]),
                ("대분류별 라인 중단 시간", tables["대분류별 라인 중단 시간"]),
            ],
        )
        write_multi_tables(
            writer,
            SHEETS["part"],
            [
                ("소분류 고장부품별 접수 건수 TOP 20", tables["소분류 고장부품별 접수 건수 TOP 20"]),
                ("소분류 고장부품별 라인 중단 시간 TOP 20", tables["소분류 고장부품별 라인 중단 시간 TOP 20"]),
            ],
        )
        write_table(writer, tables["라인 중단 시간 TOP 20 이슈 목록"], SHEETS["downtime_top"])
        write_table(writer, tables["진행중/미완료 이슈 목록"], SHEETS["open_issues"])
        write_table(writer, _quality_to_frame(analysis["quality"]), SHEETS["quality"])
        write_table(writer, analysis.get("comparison", pd.DataFrame()), SHEETS["comparison"])
        write_table(writer, repeat_issue_excel_frame(analysis.get("repeat_issues", pd.DataFrame())), SHEETS["repeat_issues"])
        write_raw_quality_sheet(writer, analysis.get("raw_quality", {}), SHEETS["raw_quality"])
        write_legacy_feature_sheets(writer, analysis.get("legacy_features", {}))
        write_quality_feedback_sheet(writer, analysis.get("quality_feedback", {}))

        for worksheet in writer.book.worksheets:
            format_sheet(worksheet)


def write_legacy_feature_sheets(writer: pd.ExcelWriter, legacy_features: dict[str, Any]) -> None:
    write_multi_tables(
        writer,
        SHEETS["claim_analysis"],
        [
            ("클레임/제조사 요약", legacy_features.get("claim_summary", pd.DataFrame())),
            ("제조사별 클레임/부품수리 접수", legacy_features.get("claim_by_manufacturer", pd.DataFrame())),
            ("유/무상 분포", legacy_features.get("claim_by_paid_type", pd.DataFrame())),
            ("분류별 클레임/부품수리 접수", legacy_features.get("claim_by_part_category", pd.DataFrame())),
            ("클레임/부품수리 상세 목록", legacy_features.get("claim_issue_list", pd.DataFrame())),
        ],
    )
    write_multi_tables(
        writer,
        SHEETS["completion_status"],
        [
            ("처리 상태 KPI", legacy_features.get("completion_kpi", pd.DataFrame())),
            ("월별 처리 상태 추이", legacy_features.get("monthly_completion_trend", pd.DataFrame())),
            ("미완료/진행 상태 분포", legacy_features.get("open_status_breakdown", pd.DataFrame())),
        ],
    )


def write_quality_feedback_sheet(writer: pd.ExcelWriter, quality_feedback: dict[str, Any]) -> None:
    write_multi_tables(
        writer,
        SHEETS["quality_feedback"],
        [
            ("품질 피드백 검토 후보 요약", quality_feedback.get("quality_feedback_summary", pd.DataFrame())),
            ("검토 후보 유형별 건수", quality_feedback.get("quality_feedback_by_candidate_type", pd.DataFrame())),
            ("클레임/제조사 입력 검토 후보", quality_feedback.get("quality_feedback_claim_candidates", pd.DataFrame())),
            ("유/무상/제조사 미입력 검토 후보", quality_feedback.get("quality_feedback_missing_field_candidates", pd.DataFrame())),
            ("부품/분류별 접수 검토 후보", quality_feedback.get("quality_feedback_part_category_candidates", pd.DataFrame())),
            ("이슈 preview", quality_feedback.get("quality_feedback_issue_preview", pd.DataFrame())),
        ],
    )
def _dict_to_frame(values: dict[str, Any]) -> pd.DataFrame:
    label_map = {
        "라인 중단 총 시간": "라인 중단 시간(입력 건 기준)",
        "라인 중단 평균 시간": "라인 중단 평균 시간(입력 건 기준)",
        "라인 중단 최대 시간": "라인 중단 최대 시간(입력 건 기준)",
    }
    return pd.DataFrame([{"항목": label_map.get(key, key), "값": value} for key, value in values.items()])


def _quality_to_frame(values: dict[str, Any]) -> pd.DataFrame:
    label_map = {
        "전체 원본 건수": "전체 원본 행 수",
        "분석 대상 건수": "선택 기간 분석 대상 행 수",
    }
    return pd.DataFrame([{"항목": label_map.get(key, key), "값": value} for key, value in values.items()])


def repeat_issue_excel_frame(dataframe: pd.DataFrame) -> pd.DataFrame:
    if dataframe.empty:
        return dataframe
    output = dataframe.copy()
    if "조건값" in output.columns and "조건값_표시" not in output.columns:
        output["조건값_표시"] = output["조건값"].map(format_repeat_condition_display)
    if "조건값" in output.columns and "조건값_품질" not in output.columns:
        output["조건값_품질"] = output["조건값"].map(evaluate_repeat_condition_quality)
    preferred = ["반복 기준", "조건값", "조건값_표시", "조건값_품질", "접수 건수", "라인 중단 총 시간", "최초 접수일", "최근 접수일", "관련 ID 목록"]
    return output.loc[:, [column for column in preferred if column in output.columns]]


def write_raw_quality_sheet(writer: pd.ExcelWriter, raw_quality: dict[str, Any], sheet_name: str) -> None:
    if not raw_quality:
        write_table(writer, pd.DataFrame([{"항목": "품질 상태", "값": "미확인"}]), sheet_name)
        return
    summary = raw_quality.get("summary", {})
    field_quality = raw_quality.get("field_quality_df", pd.DataFrame())
    conditional = raw_quality.get("conditional_field_df", pd.DataFrame())
    manual_review = raw_quality.get("manual_review_targets_df", pd.DataFrame())
    field_top = field_quality.sort_values("실제 보완 필요 건수", ascending=False).head(10) if not field_quality.empty and "실제 보완 필요 건수" in field_quality.columns else field_quality
    conditional_count = (
        conditional.groupby(["업무유형", "제외 필드"], dropna=False).size().reset_index(name="필드 플래그 수")
        if not conditional.empty and {"업무유형", "제외 필드"}.issubset(conditional.columns)
        else pd.DataFrame(columns=["업무유형", "제외 필드", "필드 플래그 수"])
    )
    write_multi_tables(
        writer,
        sheet_name,
        [
            ("Raw 품질 요약", _raw_quality_summary_frame(raw_quality)),
            ("필드별 실제 보완 필요 TOP 10", field_top),
            ("조건부 제외 유/무상, 제조사 count", conditional_count),
            ("수동 보완 TOP 20", manual_review.head(20)),
        ],
    )


def _raw_quality_summary_frame(raw_quality: dict[str, Any]) -> pd.DataFrame:
    summary = raw_quality.get("summary", {})
    rows = [
        ("품질 상태", raw_quality.get("status", "")),
        ("전체 원본 행 수", summary.get("전체 건수", 0)),
        ("선택 기간 분석 대상 행 수", summary.get("분석 대상 건수", 0)),
        ("보완 필요 행 수", summary.get("보완 필요 건수", 0)),
        ("우선순위 상", summary.get("우선순위 상", 0)),
        ("우선순위 중", summary.get("우선순위 중", 0)),
        ("우선순위 하", summary.get("우선순위 하", 0)),
        ("정상 공백/조건부 제외 필드 플래그 수", summary.get("조건부 제외 건수", 0)),
    ]
    return pd.DataFrame([{"항목": key, "값": value} for key, value in rows])


def write_table(writer: pd.ExcelWriter, dataframe: pd.DataFrame, sheet_name: str, startrow: int = 0) -> int:
    output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "데이터가 없습니다."}])
    output.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow)
    return startrow + len(output) + 3


def write_multi_tables(writer: pd.ExcelWriter, sheet_name: str, tables: list[tuple[str, pd.DataFrame]]) -> None:
    workbook = writer.book
    worksheet = workbook.create_sheet(sheet_name)
    writer.sheets[sheet_name] = worksheet
    startrow = 0
    for title, dataframe in tables:
        worksheet.cell(row=startrow + 1, column=1, value=title)
        worksheet.cell(row=startrow + 1, column=1).font = Font(bold=True)
        startrow = write_table(writer, dataframe, sheet_name, startrow=startrow + 1)


def format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and len(cell.value) > 24:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")
            if hasattr(cell.value, "strftime"):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, float):
                header = worksheet.cell(row=1, column=cell.column).value
                cell.number_format = "0.0%" if isinstance(header, str) and "율" in header and abs(cell.value) <= 1 else "#,##0.0"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"

    for column_cells in worksheet.columns:
        max_length = 10
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        worksheet.column_dimensions[column_letter].width = max_length + 2
