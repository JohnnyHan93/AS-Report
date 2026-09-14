from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_manual_review_workbook(quality_result: dict[str, Any], output_dir: str | Path) -> str:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    workbook_path = output_path / f"raw_data_manual_review_targets_{timestamp}.xlsx"

    summary = _summary_frame(quality_result.get("summary", {}), quality_result.get("status", ""))
    manual_review = quality_result.get("manual_review_targets_df", pd.DataFrame())
    field_quality = quality_result.get("field_quality_df", pd.DataFrame())
    conditional = quality_result.get("conditional_field_df", pd.DataFrame())

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        _usage_frame().to_excel(writer, sheet_name="사용방법", index=False)
        summary.to_excel(writer, sheet_name="요약", index=False)
        _write_table(writer, manual_review, "보완대상_검색용")
        _write_table(writer, manual_review.head(60), "우선검색_TOP")
        _write_table(writer, field_quality, "필드별_보완건수")
        _write_table(writer, conditional, "조건부_유무상제조사")
        _search_rule_frame().to_excel(writer, sheet_name="검색기준", index=False)

        for worksheet in writer.book.worksheets:
            _format_sheet(worksheet)

    return str(workbook_path)


def _write_table(writer: pd.ExcelWriter, dataframe: pd.DataFrame, sheet_name: str) -> None:
    output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "대상 데이터가 없습니다."}])
    output.to_excel(writer, sheet_name=sheet_name, index=False)


def _usage_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"항목": "목적", "내용": "Raw 데이터에서 보완 필요 후보를 찾아 그룹웨어 또는 담당자 확인에 사용할 수 있도록 정리합니다."},
            {"항목": "주의", "내용": "자동추정 값은 확정값이 아니며 원본 CSV를 수정하지 않습니다."},
            {"항목": "검색", "내용": "추천 검색방법, 접수일, 접수내용 요약 검색어를 조합해 원본 접수 건을 확인하세요."},
            {"항목": "결과", "내용": "우선순위 상 항목부터 확인하고, 필요 시 원본 시스템에서 값을 보완하세요."},
        ]
    )


def _search_rule_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"우선순위": 1, "검색 기준": "접수일 + 접수내용 요약 검색어"},
            {"우선순위": 2, "검색 기준": "접수일 + 고객사 + 업무유형"},
            {"우선순위": 3, "검색 기준": "접수일 + 접수사원 + 일부 키워드"},
            {"우선순위": 4, "검색 기준": "접수일 범위 + 고객사 + 로보트/설비 키워드"},
        ]
    )


def _summary_frame(values: dict[str, Any], status: str) -> pd.DataFrame:
    rows = [
        ("품질 상태", status),
        ("Raw 데이터 파일", values.get("Raw 데이터 파일", "")),
        ("접수일 시작", values.get("접수일 시작", "")),
        ("접수일 종료", values.get("접수일 종료", "")),
        ("전체 원본 행 수", values.get("전체 건수", 0)),
        ("선택 기간 분석 대상 행 수", values.get("분석 대상 건수", 0)),
        ("보완 필요 행 수", values.get("보완 필요 건수", 0)),
        ("우선순위 상", values.get("우선순위 상", 0)),
        ("우선순위 중", values.get("우선순위 중", 0)),
        ("우선순위 하", values.get("우선순위 하", 0)),
        ("정상 공백/조건부 제외 필드 플래그 수", values.get("조건부 제외 건수", 0)),
        ("오류값 건수", values.get("오류값 건수", 0)),
    ]
    return pd.DataFrame([{"항목": key, "값": value} for key, value in rows])


def _format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows():
        for cell in row:
            if hasattr(cell.value, "strftime"):
                cell.number_format = "yyyy-mm-dd"
            if isinstance(cell.value, str) and len(cell.value) > 20:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 10
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 55))
        worksheet.column_dimensions[column_letter].width = max_length + 2
