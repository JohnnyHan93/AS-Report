from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from as_report.output import build_timestamped_filename, get_run_timestamp

from .standard_name_mapper import FIELD_TO_FILE, MAPPING_COLUMNS, TARGET_FIELDS


def export_unknown_standard_names_workbook(unknown_df: pd.DataFrame, output_dir: str | Path, timestamp: str | None = None) -> str:
    """Write unknown standard names to a review workbook."""
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    run_timestamp = timestamp or get_run_timestamp()
    workbook_path = output_path / build_timestamped_filename("unknown_standard_names", ".xlsx", run_timestamp)

    unknown = _unknown_frame(unknown_df)
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        _usage_frame().to_excel(writer, sheet_name="사용방법", index=False)
        _summary_frame(unknown).to_excel(writer, sheet_name="요약", index=False)
        unknown.to_excel(writer, sheet_name="Unknown_명칭목록", index=False)
        _field_count_frame(unknown).to_excel(writer, sheet_name="필드별_건수", index=False)
        _guide_frame().to_excel(writer, sheet_name="매핑CSV_작성가이드", index=False)
        for worksheet in writer.book.worksheets:
            _format_sheet(worksheet)
    return str(workbook_path)


def export_standard_name_preview_csv(preview_df: pd.DataFrame, output_dir: str | Path, timestamp: str | None = None) -> str:
    """Write standard-name preview columns to a CSV file."""
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    csv_path = output_path / build_timestamped_filename("standard_name_preview", ".csv", timestamp or get_run_timestamp())
    preview_df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    return str(csv_path)


def _usage_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"순서": 1, "내용": "Unknown_명칭목록 시트에서 원본값을 확인한다."},
            {"순서": 2, "내용": r"해당 값을 config\standard_names 아래 매핑 CSV에 입력한다."},
            {"순서": 3, "내용": "raw_name에는 원본값, standard_name에는 표준값을 입력한다."},
            {"순서": 4, "내용": "enabled는 TRUE로 입력한다."},
            {"순서": 5, "내용": "원본 Raw CSV는 수정하지 않는다."},
        ]
    )


def _summary_frame(unknown: pd.DataFrame) -> pd.DataFrame:
    if unknown.empty:
        return pd.DataFrame([{"항목": "Unknown 명칭 수", "값": 0}])
    return pd.DataFrame(
        [
            {"항목": "Unknown 명칭 수", "값": int(len(unknown))},
            {"항목": "Unknown 발생 건수", "값": int(unknown["count"].sum()) if "count" in unknown.columns else 0},
            {"항목": "대상 필드 수", "값": int(unknown["field_name"].nunique()) if "field_name" in unknown.columns else 0},
        ]
    )


def _unknown_frame(unknown_df: pd.DataFrame) -> pd.DataFrame:
    expected_columns = [
        "field_name",
        "raw_name",
        "count",
        "example_접수일",
        "example_접수내용",
        "recommended_mapping_file",
        "note",
    ]
    if unknown_df.empty:
        return pd.DataFrame(columns=expected_columns)
    output = unknown_df.copy()
    for column in expected_columns:
        if column not in output.columns:
            output[column] = ""
    return output[expected_columns]


def _field_count_frame(unknown: pd.DataFrame) -> pd.DataFrame:
    if unknown.empty:
        return pd.DataFrame(columns=["field_name", "unknown_count", "unique_unknown_count"])
    grouped = unknown.groupby("field_name", dropna=False)
    return (
        grouped.agg(unknown_count=("count", "sum"), unique_unknown_count=("raw_name", "nunique"))
        .reset_index()
        .sort_values(["unknown_count", "field_name"], ascending=[False, True])
    )


def _guide_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "field_name": field_name,
                "mapping_file": FIELD_TO_FILE.get(field_name, ""),
                "required_columns": ", ".join(MAPPING_COLUMNS),
                "description": "raw_name 원본값과 standard_name 표준값을 입력하고 enabled를 TRUE로 설정합니다.",
            }
            for field_name in TARGET_FIELDS
        ]
    )


def _format_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 25)
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 10
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        worksheet.column_dimensions[column_letter].width = max_length + 2
