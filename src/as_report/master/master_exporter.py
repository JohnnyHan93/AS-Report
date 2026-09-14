from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from as_report.output import build_timestamped_filename, get_run_timestamp


def export_master_analysis_outputs(master_result: dict[str, Any], output_dir: str | Path, timestamp: str | None = None) -> dict[str, Path]:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    run_timestamp = timestamp or get_run_timestamp()
    paths = {
        "customer_robot_master_quality": output_path / build_timestamped_filename("customer_robot_master_quality", ".xlsx", run_timestamp),
        "failure_part_master_quality": output_path / build_timestamped_filename("failure_part_master_quality", ".xlsx", run_timestamp),
        "as_robot_match_preview": output_path / build_timestamped_filename("as_robot_match_preview", ".csv", run_timestamp),
        "unmatched_as_issues": output_path / build_timestamped_filename("unmatched_as_issues", ".xlsx", run_timestamp),
        "customer_as_rate_by_robot_count": output_path / build_timestamped_filename("customer_as_rate_by_robot_count", ".xlsx", run_timestamp),
        "customer_as_rate_summary": output_path / build_timestamped_filename("customer_as_rate_summary", ".md", run_timestamp),
        "failure_part_category_validation": output_path / build_timestamped_filename("failure_part_category_validation", ".csv", run_timestamp),
        "failure_part_trend": output_path / build_timestamped_filename("failure_part_trend", ".xlsx", run_timestamp),
    }

    _write_quality_workbook(paths["customer_robot_master_quality"], "Customer Robot Master 품질", master_result.get("customer_robot_quality", {}))
    _write_quality_workbook(paths["failure_part_master_quality"], "Failure Part Master 품질", master_result.get("failure_part_quality", {}))
    _to_csv(master_result.get("robot_match_preview_df", pd.DataFrame()), paths["as_robot_match_preview"])
    _write_table_workbook(paths["unmatched_as_issues"], "미매칭_AS", master_result.get("unmatched_as_issues_df", pd.DataFrame()))
    _write_table_workbook(paths["customer_as_rate_by_robot_count"], "고객사별_AS접수율", master_result.get("customer_as_rate_df", pd.DataFrame()))
    _write_customer_rate_summary(paths["customer_as_rate_summary"], master_result)
    _to_csv(master_result.get("category_validation_df", pd.DataFrame()), paths["failure_part_category_validation"])
    _write_table_workbook(paths["failure_part_trend"], "고장부품_추세", master_result.get("failure_part_trend_df", pd.DataFrame()))
    return paths


def _write_quality_workbook(path: Path, title: str, quality: dict[str, Any]) -> None:
    summary = quality.get("summary", {})
    field_quality = quality.get("field_quality_df", pd.DataFrame())
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _dict_frame(summary).to_excel(writer, sheet_name="요약", index=False)
        _write_table(writer, field_quality, "필드별_품질")
        _format_workbook(writer, title)


def _write_table_workbook(path: Path, sheet_name: str, dataframe: pd.DataFrame) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_table(writer, dataframe, sheet_name)
        _format_workbook(writer, sheet_name)


def _write_customer_rate_summary(path: Path, master_result: dict[str, Any]) -> None:
    summary = master_result.get("summary", {})
    customer_rate = master_result.get("customer_as_rate_df", pd.DataFrame())
    top_rows = customer_rate.head(10) if not customer_rate.empty else pd.DataFrame()
    lines = [
        "# Customer AS Rate Summary",
        "",
        "## 분석 요약",
        f"- AS 분석 대상 행 수: {summary.get('AS 분석 대상 행 수', 0):,}",
        f"- 로봇 master 행 수: {summary.get('로봇 master 행 수', 0):,}",
        f"- 고장부품 master 행 수: {summary.get('고장부품 master 행 수', 0):,}",
        f"- 로봇 master 매칭 행 수: {summary.get('로봇 master 매칭 행 수', 0):,}",
        f"- 로봇 master 미매칭 행 수: {summary.get('로봇 master 미매칭 행 수', 0):,}",
        f"- 고장성 AS 행 수: {summary.get('고장성 AS 행 수', 0):,}",
        "",
        "## 해석 주의",
        "- 본 요약은 설치 로봇 수 대비 AS 접수율을 계산한 참고 지표입니다.",
        "- 고장성 AS는 고장부품 master와 매칭되는 건만 포함합니다.",
        "- 매칭 결과는 확정값이 아니라 검토용 preview입니다.",
        "",
        "## 고객사별 TOP 10",
        _markdown_table(top_rows),
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _to_csv(dataframe: pd.DataFrame, path: Path) -> None:
    output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "데이터가 없습니다."}])
    output.to_csv(path, index=False, encoding="utf-8-sig")


def _write_table(writer: pd.ExcelWriter, dataframe: pd.DataFrame, sheet_name: str) -> None:
    output = dataframe if not dataframe.empty else pd.DataFrame([{"안내": "데이터가 없습니다."}])
    output.to_excel(writer, sheet_name=sheet_name, index=False)


def _dict_frame(values: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{"항목": key, "값": value} for key, value in values.items()])


def _format_workbook(writer: pd.ExcelWriter, title: str) -> None:
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        header_fill = PatternFill("solid", fgColor="E2F0D9")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for row in worksheet.iter_rows():
            for cell in row:
                if hasattr(cell.value, "strftime"):
                    cell.number_format = "yyyy-mm-dd"
                if isinstance(cell.value, str) and len(cell.value) > 24:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(vertical="top")
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = 10
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 60))
            worksheet.column_dimensions[column_letter].width = max_length + 2
    writer.book.properties.title = title


def _markdown_table(dataframe: pd.DataFrame) -> str:
    if dataframe.empty:
        return "- 데이터가 없습니다."
    columns = dataframe.columns.tolist()
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for _, row in dataframe.iterrows():
        lines.append("| " + " | ".join(str(row.get(column, "")) for column in columns) + " |")
    return "\n".join(lines)
